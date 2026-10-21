from __future__ import annotations

import asyncio
from pathlib import Path

import pymupdf
from docx import Document
from openpyxl import Workbook, load_workbook
from pptx import Presentation
from pptx.util import Inches

from ee_pipeline.chatt_client import chat_completions_url, validate_chatt_config
from ee_pipeline.config import PipelineConfig
from ee_pipeline.models import (
    ClassificationResult,
    DocumentBlock,
    EECategory,
    FileScreenResult,
    KeywordMatch,
    TranslationResult,
)
from ee_pipeline.modules.module1_loader import discover_files
from ee_pipeline.modules.module2_5_prescreen import _extract_keywords
from ee_pipeline.modules.module2_5_prescreen import run as run_prescreen
from ee_pipeline.modules.module2_parser import _detect_heading, parse_document
from ee_pipeline.modules.module3_chunker import chunk_blocks
from ee_pipeline.modules.module4_5_llm_input import prepare_llm_inputs
from ee_pipeline.modules.module4_segmenter import segment_chunks
from ee_pipeline.modules.module5_classifier import (
    _classify_single_async,
    _parse_llm_response,
    _validate_classification_payload,
)
from ee_pipeline.modules.module6_translate import _translate_single_async, _translation_result
from ee_pipeline.modules.module7_aggregator import aggregate
from ee_pipeline.modules.module8_report import generate_report


def test_m1_manifest_ids_are_deterministic_and_unique(tmp_path: Path) -> None:
    (tmp_path / "same.pdf").write_bytes(b"")
    nested = tmp_path / "nested"
    nested.mkdir()
    (nested / "same.docx").write_bytes(b"")

    manifest = discover_files(tmp_path)

    assert [entry["doc_id"] for entry in manifest] == ["same", "same__2"]
    assert all(Path(entry["path"]).is_absolute() for entry in manifest)


def test_m2_docx_parser_preserves_heading_provenance(tmp_path: Path) -> None:
    path = tmp_path / "requirements.docx"
    document = Document()
    document.add_heading("1 System", level=1)
    document.add_paragraph("The sensor shall provide a valid signal.")
    document.save(path)

    blocks = parse_document(str(path), "requirements", "docx")

    assert len(blocks) == 1
    assert blocks[0].heading == "1 System"
    assert blocks[0].heading_index == 1
    assert "shall" in blocks[0].content


def test_m2_does_not_treat_table_fragment_as_heading() -> None:
    assert _detect_heading("30 %", font_size=10, body_font_size=9) is None
    assert _detect_heading("2.1  General", font_size=12, body_font_size=10) == {"level": 3}
    assert _detect_heading("1\nScope of this Document", font_size=14, body_font_size=10) == {"level": 2}


def test_m2_pdf_metadata_state_does_not_leak_to_later_pages(tmp_path: Path) -> None:
    path = tmp_path / "datasheet.pdf"
    document = pymupdf.open()
    contents = document.new_page()
    contents.insert_text((72, 72), "Table of Contents", fontsize=16)
    contents.insert_text((72, 100), "1 Overview ........ 3", fontsize=10)
    body = document.new_page()
    body.insert_text((72, 72), "The controller shall detect a valid signal.", fontsize=10)
    document.save(path)
    document.close()

    blocks = parse_document(str(path), "datasheet", "pdf")
    requirement = next(block for block in blocks if "shall detect" in block.content)

    assert requirement.section_role != "boilerplate"


def test_m2_metadata_keyword_inside_body_sentence_is_not_boilerplate(tmp_path: Path) -> None:
    path = tmp_path / "wake_up.pdf"
    document = pymupdf.open()
    page = document.new_page()
    page.insert_text((72, 72), "7.4 Remote wake-up", fontsize=14)
    page.insert_text(
        (72, 105),
        "The wake-up pattern contents a dominant signal within the maximum wake-up time.",
        fontsize=10,
    )
    document.save(path)
    document.close()

    blocks = parse_document(str(path), "wake_up", "pdf")
    requirement = next(block for block in blocks if "pattern contents" in block.content)

    assert requirement.section_role != "boilerplate"


def test_m2_docx_parser_preserves_table_rows(tmp_path: Path) -> None:
    path = tmp_path / "requirements_with_table.docx"
    document = Document()
    document.add_heading("1 Requirements", level=1)
    table = document.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Requirement ID"
    table.rows[0].cells[1].text = "Requirement"
    table.rows[1].cells[0].text = "SYS-001"
    table.rows[1].cells[1].text = "The controller shall report faults."
    document.save(path)

    blocks = parse_document(str(path), "requirements", "docx")

    assert len(blocks) == 1
    assert blocks[0].block_type == "table"
    assert blocks[0].record_id == "SYS-001"
    assert "Requirement: The controller shall report faults." in blocks[0].content


def test_m2_pptx_preserves_title_only_slides_and_table_cells(tmp_path: Path) -> None:
    path = tmp_path / "requirements.pptx"
    presentation = Presentation()
    title_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    title_slide.shapes.title.text = "Architecture Overview"
    table_slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    table_slide.shapes.title.text = "Interface Requirements"
    table = table_slide.shapes.add_table(2, 2, Inches(1), Inches(2), Inches(7), Inches(2)).table
    table.cell(0, 0).text = "ID"
    table.cell(0, 1).text = "Requirement"
    table.cell(1, 0).text = "IF-001"
    table.cell(1, 1).text = "The interface shall detect timeout."
    presentation.save(path)

    blocks = parse_document(str(path), "requirements", "pptx")

    assert {block.page_number for block in blocks} == {1, 2}
    assert any(block.page_number == 1 and "Architecture Overview" in block.content for block in blocks)
    assert any(
        block.block_type == "table" and "Requirement: The interface shall detect timeout." in block.content
        for block in blocks
    )


def test_m2_xlsx_parser_keeps_requirement_rows_atomic(tmp_path: Path) -> None:
    path = tmp_path / "requirements.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Requirements"
    sheet.append(["Requirement ID", "Requirement", "Priority"])
    sheet.append(["SYS-INP-001", "The system shall receive input.", "High"])
    sheet.append(["SYS-INP-002", "The system shall validate input.", "High"])
    workbook.save(path)

    blocks = parse_document(str(path), "requirements", "xlsx")

    assert [block.record_id for block in blocks] == ["SYS-INP-001", "SYS-INP-002"]
    assert all(block.block_type == "table" for block in blocks)
    assert "Requirement: The system shall receive input." in blocks[0].content


def test_m2_xlsx_identifies_support_sheet_by_schema(tmp_path: Path) -> None:
    path = tmp_path / "mixed.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Checks"
    sheet.append(["Requirement ID", "Verification Method", "Test Case ID", "Evidence Status"])
    sheet.append(["SYS-001", "Test", "TC-001", "Required evidence missing"])
    workbook.save(path)

    blocks = parse_document(str(path), "mixed", "xlsx")

    assert len(blocks) == 1
    assert blocks[0].block_type == "table"
    assert blocks[0].section_role == "support"


def test_m25_matches_terms_without_substring_false_positives(tmp_path: Path) -> None:
    config = PipelineConfig(
        input_dir=tmp_path,
        output_dir=tmp_path,
        _run_dir=tmp_path / "run",
        keywords={"primary": {"network": ["CAN"]}, "secondary": {"sensor": ["sensor"]}},
        min_primary_keywords=1,
    )
    block = DocumentBlock(
        doc_id="doc",
        block_id="doc_b1",
        content="CAN is valid; candy is not. The sensor is measured.",
        source_file="doc.docx",
    )

    matches = _extract_keywords(block.content, config.keywords)
    passed, skipped = run_prescreen(config, [block])

    assert [(match.keyword, match.count) for match in matches] == [("CAN", 1), ("sensor", 1)]
    assert len(passed) == 1 and not skipped


def test_m3_does_not_cross_document_boundaries() -> None:
    blocks = [
        DocumentBlock(doc_id="a", block_id="a1", heading="A", heading_level=1, content="A", source_file="a.pdf"),
        DocumentBlock(doc_id="b", block_id="b1", heading="B", heading_level=1, content="B", source_file="b.pdf"),
    ]

    chunks = chunk_blocks(blocks)

    assert len(chunks) == 2
    assert [chunk.blocks[0].doc_id for chunk in chunks] == ["a", "b"]


def test_m4_batch_and_segment_ids_are_unique() -> None:
    blocks = [
        DocumentBlock(
            doc_id="doc",
            block_id="b1",
            heading="Sensors",
            heading_level=1,
            content="The sensor shall provide a valid signal.\n\nThe sensor must detect faults.",
            source_file="doc.pdf",
            heading_index=1,
            source_start_line=10,
            source_end_line=12,
        )
    ]

    batches = segment_chunks(chunk_blocks(blocks))
    segment_ids = [source["segment_id"] for batch in batches for source in batch["source_segments"]]

    assert len({batch["batch_id"] for batch in batches}) == len(batches)
    assert len(set(segment_ids)) == len(segment_ids)
    assert batches[0]["source_segments"][0]["source_file"] == "doc.pdf"


def test_m4_classifies_blocks_individually_inside_mixed_chunk() -> None:
    blocks = [
        DocumentBlock(
            doc_id="doc",
            block_id="p1",
            heading="7.2 Wake-up",
            heading_level=2,
            content="The device shall detect a valid wake-up pattern.",
            source_file="doc.pdf",
            block_type="paragraph",
        ),
        DocumentBlock(
            doc_id="doc",
            block_id="f1",
            heading="7.2 Wake-up",
            heading_level=2,
            content="Figure 3. Wake-up timing",
            source_file="doc.pdf",
            block_type="figure",
        ),
    ]

    batches = segment_chunks(chunk_blocks(blocks), include_skipped=True)

    assert any(not batch["skip_candidate"] and "shall detect" in batch["merged_text"] for batch in batches)
    assert any(batch["skip_candidate"] and batch["candidate_type"] == "figure" for batch in batches)


def test_m4_does_not_treat_contents_verb_as_boilerplate() -> None:
    blocks = [
        DocumentBlock(
            doc_id="doc",
            block_id="wake_up",
            heading="7.4 Remote wake-up",
            heading_level=3,
            content=(
                "Within the maximum wake-up time, the wake-up pattern contents a dominant signal "
                "followed by a recessive signal."
            ),
            source_file="doc.pdf",
            block_type="paragraph",
            section_role="content",
        )
    ]

    batches = segment_chunks(chunk_blocks(blocks), include_skipped=True)

    assert len(batches) == 1
    assert batches[0]["skip_candidate"] is False


def test_m4_separates_section_role_from_physical_type() -> None:
    blocks = [
        DocumentBlock(
            doc_id="doc",
            block_id="front_req",
            heading="Front matter",
            content="The controller shall preserve diagnostic data.",
            source_file="doc.pdf",
            block_type="paragraph",
            section_role="metadata",
        ),
        DocumentBlock(
            doc_id="doc",
            block_id="verification",
            heading="Verification data",
            content="Required evidence missing for TC-001.",
            source_file="doc.xlsx",
            block_type="table",
            section_role="support",
            source_unit_id="row-1",
        ),
    ]

    batches = segment_chunks(chunk_blocks(blocks), include_skipped=True)

    assert any(not batch["skip_candidate"] and "shall preserve" in batch["merged_text"] for batch in batches)
    assert any(
        batch["skip_candidate"]
        and batch["source_segments"][0]["skip_reason"] == "supporting_material"
        for batch in batches
    )


def test_m4_does_not_merge_atomic_source_units() -> None:
    blocks = [
        DocumentBlock(
            doc_id="slides",
            block_id="slide_1",
            heading="Repeated title",
            content="First slide body.",
            source_file="slides.pptx",
            source_unit_id="slide-1",
            page_number=1,
        ),
        DocumentBlock(
            doc_id="slides",
            block_id="slide_2",
            heading="Repeated title",
            content="Second slide body.",
            source_file="slides.pptx",
            source_unit_id="slide-2",
            page_number=2,
        ),
    ]

    batches = segment_chunks(chunk_blocks(blocks))

    assert len(batches) == 2
    assert [batch["page_numbers"] for batch in batches] == [[1], [2]]


def test_m45_creates_stable_cur_id_and_clean_requirement_text() -> None:
    records = prepare_llm_inputs(
        [
            {
                "document_id": "customer",
                "batch_id": "batch_0001",
                "candidate_type": "spec_table",
                "heading": "Requirements",
                "source_file": "customer.xlsx",
                "page_numbers": [1],
                "record_id": "SYS-001",
                "source_segments": [
                    {
                        "segment_text": (
                            "Requirement ID: SYS-001 | Requirement: The ECU shall detect a fault. "
                            "| Status: Draft"
                        ),
                        "record_id": "SYS-001",
                    }
                ],
            }
        ]
    )

    assert records[0].cur_id == "SYS-001"
    assert records[0].original_text == "The ECU shall detect a fault."
    assert "Customer requirement ID: SYS-001" in records[0].llm_text


def test_m45_splits_generic_specification_tables_into_atomic_rows() -> None:
    records = prepare_llm_inputs(
        [
            {
                "document_id": "device",
                "batch_id": "batch_0001",
                "candidate_type": "spec_table",
                "heading": "Electrical characteristics",
                "source_file": "device.pdf",
                "page_numbers": [3],
                "table_title": "Table 1",
                "source_segments": [
                    {"segment_text": "Table 1. Electrical characteristics", "page_number": 3},
                    {
                        "segment_text": "Symbol: VCC | Parameter: supply voltage | Min: 4.5 | Max: 5.5 | Unit: V",
                        "page_number": 3,
                    },
                    {
                        "segment_text": "Symbol: ICC | Parameter: supply current | Max: 10 | Unit: mA",
                        "page_number": 3,
                    },
                ],
            }
        ]
    )

    assert len(records) == 2
    assert {record.parent_batch_id for record in records} == {"batch_0001"}
    assert len({record.cur_id for record in records}) == 2
    assert records[0].original_text.startswith("Symbol: VCC")


def test_m5_response_parser_handles_fenced_json() -> None:
    payload = _parse_llm_response('prefix\n```json\n{"is_requirement": true, "is_ee_related": false}\n```')

    assert payload == {"is_requirement": True, "is_ee_related": False}


def test_m5_rejects_ambiguous_or_incomplete_classification() -> None:
    for payload in (
        {"is_requirement": "maybe", "is_ee_related": False},
        {"is_requirement": True, "is_ee_related": True},
        {"is_requirement": True, "is_ee_related": True, "ee_category": "not-a-category"},
    ):
        try:
            _validate_classification_payload(payload)
        except ValueError:
            pass
        else:
            raise AssertionError(f"Invalid ChatT payload must be rejected: {payload}")


def test_m6_marks_short_translation_for_review() -> None:
    classification = ClassificationResult(
        block_id="b1",
        doc_id="doc",
        is_requirement=True,
        is_ee_related=True,
        original_content="The sensor shall work.",
    )

    result = _translation_result(classification, "짧음", True)

    assert result.status == "REVIEW_NEEDED"


def test_chatt_adapters_keep_m5_m6_contracts(tmp_path: Path) -> None:
    class FakeResponse:
        def __init__(self, content: str):
            self.content = content

        def raise_for_status(self):
            return None

        def json(self):
            return {"choices": [{"message": {"content": self.content}}]}

    class FakeClient:
        def __init__(self, content: str):
            self.content = content

        async def post(self, url, *, json, headers):
            assert url == "https://chatt.example/v1/chat/completions"
            assert json["model"] == "ChatT-test"
            assert json["messages"][1]["content"]
            assert headers["Authorization"] == "Bearer test"
            return FakeResponse(self.content)

    config = PipelineConfig(
        output_dir=tmp_path,
        _run_dir=tmp_path / "run",
        chatt_api_key="test",
        chatt_base_url="https://chatt.example",
        chatt_model="ChatT-test",
    )

    async def call_adapters():
        classification = await _classify_single_async(
            "The sensor shall provide a signal.",
            client=FakeClient(
                '{"is_requirement": true, "is_ee_related": true, '
                '"ee_category": "Sensor", "reasoning": "sensor"}'
            ),
            config=config,
        )
        translation = await _translate_single_async(
            "The sensor shall provide a signal.",
            client=FakeClient("센서는 유효한 신호를 제공해야 한다."),
            config=config,
        )
        return classification, translation

    classification, translation = asyncio.run(call_adapters())
    assert classification["ee_category"] == "Sensor"
    assert translation.startswith("센서")


def test_chatt_configuration_and_endpoint_validation(tmp_path: Path) -> None:
    assert chat_completions_url("https://chatt.example") == "https://chatt.example/v1/chat/completions"
    assert chat_completions_url("https://chatt.example/v1") == "https://chatt.example/v1/chat/completions"
    assert (
        chat_completions_url("https://chatt.example/v1/chat/completions")
        == "https://chatt.example/v1/chat/completions"
    )
    assert (
        chat_completions_url("https://chatt.example/custom/chat/completions")
        == "https://chatt.example/custom/chat/completions"
    )

    config = PipelineConfig(output_dir=tmp_path, _run_dir=tmp_path, chatt_api_key="", chatt_base_url="", chatt_model="")
    try:
        validate_chatt_config(config)
    except ValueError as error:
        message = str(error)
    else:
        raise AssertionError("Missing ChatT settings must fail before an API request")

    assert "CHATT_API_KEY" in message
    assert "CHATT_BASE_URL" in message
    assert "CHATT_MODEL" in message


def test_m7_and_m8_preserve_provenance_and_create_workbook(tmp_path: Path) -> None:
    config = PipelineConfig(output_dir=tmp_path, _run_dir=tmp_path)
    classification = ClassificationResult(
        block_id="batch_0001",
        doc_id="doc",
        cur_id="SEN-001",
        is_requirement=True,
        is_ee_related=True,
        ee_category=EECategory.SENSOR,
        original_content="The sensor shall provide a valid signal.",
    )
    translation = TranslationResult(
        block_id="batch_0001",
        doc_id="doc",
        cur_id="SEN-001",
        original_text=classification.original_content,
        translated_text="센서는 유효한 신호를 제공해야 한다.",
    )
    records = aggregate(
        [classification],
        [translation],
        {
            "batch_0001": {
                "heading": "Sensors",
                "cur_id": "SEN-001",
                "source_file": "doc.pdf",
                "page_number": 3,
                "source_start_line": 12,
                "matched_keywords": ["sensor"],
            }
        },
    )
    report = generate_report(
        config,
        [{"doc_id": "doc", "path": "/tmp/doc.pdf", "doc_type": "pdf"}],
        [
            FileScreenResult(
                doc_id="doc",
                source_file="doc.pdf",
                matched_keywords=[KeywordMatch(keyword="sensor", category="sensor")],
                total_matches=1,
                decision="PASS",
            )
        ],
        [classification],
        records,
    )

    workbook = load_workbook(report, read_only=False, data_only=False)
    assert records[0].source_file == "doc.pdf"
    assert records[0].matched_keywords == ["sensor"]
    assert records[0].cur_id == "SEN-001"
    assert workbook.sheetnames == ["문서 요약", "전장 요구사항"]
    assert workbook["문서 요약"]["E6"].value == 1
    assert workbook["문서 요약"]["F6"].value == 1
    assert workbook["전장 요구사항"]["B6"].value == "SEN-001"


def test_m7_rejects_missing_or_unusable_translation() -> None:
    classification = ClassificationResult(
        block_id="batch_0001",
        doc_id="doc",
        is_requirement=True,
        is_ee_related=True,
        ee_category=EECategory.SENSOR,
        original_content="The sensor shall provide a valid signal.",
    )

    try:
        aggregate([classification], [])
    except ValueError as error:
        assert "missing translations" in str(error)
    else:
        raise AssertionError("M7 must reject a missing E/E requirement translation")

    unusable = TranslationResult(
        block_id="batch_0001",
        doc_id="doc",
        original_text=classification.original_content,
        translated_text="",
        status="REVIEW_NEEDED",
    )
    try:
        aggregate([classification], [unusable])
    except ValueError as error:
        assert "unusable translations" in str(error)
    else:
        raise AssertionError("M7 must reject an unusable E/E requirement translation")


def test_m8_rejects_summary_detail_mismatch(tmp_path: Path) -> None:
    config = PipelineConfig(output_dir=tmp_path, _run_dir=tmp_path)
    classification = ClassificationResult(
        block_id="batch_0001",
        doc_id="doc",
        is_requirement=True,
        is_ee_related=True,
        ee_category=EECategory.SENSOR,
        original_content="The sensor shall provide a valid signal.",
    )

    try:
        generate_report(
            config,
            [{"doc_id": "doc", "path": "doc.pdf", "doc_type": "pdf"}],
            [],
            [classification],
            [],
        )
    except ValueError as error:
        assert "summary/detail inputs are inconsistent" in str(error)
    else:
        raise AssertionError("M8 must reject a workbook whose summary and detail rows disagree")
