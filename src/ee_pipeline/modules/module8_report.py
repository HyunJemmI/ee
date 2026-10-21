"""M8: generate the two-sheet customer requirement review workbook."""

from __future__ import annotations

from collections import Counter, defaultdict
from collections.abc import Iterable
from datetime import datetime
from math import ceil
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..config import PipelineConfig
from ..models import ClassificationResult, FileScreenResult, SpecRecord

NAVY = "17365D"
BLUE = "2F75B5"
LIGHT_BLUE = "DCE6F1"
LIGHT_GREEN = "E2F0D9"
LIGHT_RED = "FCE4D6"
WHITE = "FFFFFF"
BODY_FONT = Font(name="Arial", size=10, color="222222")
HEADER_FONT = Font(name="Arial", bold=True, size=10, color=WHITE)
THIN_GRAY = Side(style="thin", color="D9E1F2")
BOTTOM_BORDER = Border(bottom=THIN_GRAY)

SUMMARY_HEADERS = [
    "doc_id",
    "파일명",
    "문서 형식",
    "전장품 관련 여부",
    "전체 요구사항 수",
    "전장품 관련 요구사항 수",
    "전장품 요구사항 비율",
]
DETAIL_HEADERS = ["doc_id", "cur_id", "원문", "한글 번역", "비고", "전장품 종류"]

CATEGORY_KO = {
    "Sensor": "센서",
    "ECU & Software": "ECU 및 소프트웨어",
    "Actuator & Motor": "액추에이터 및 모터",
    "Power Electronics": "전력전자",
    "Power & Wiring": "전원 및 배선",
    "Communication & Diagnostics": "통신 및 진단",
    "EMC & Electrical Environment": "EMC 및 전기 환경",
    "Other": "기타",
}


def _validate_report_inputs(
    classifications: list[ClassificationResult],
    spec_records: list[SpecRecord],
) -> None:
    """Prevent a summary/detail mismatch from producing a misleading workbook."""
    expected = Counter(
        item.block_id
        for item in classifications
        if item.is_requirement and item.is_ee_related
    )
    actual = Counter(record.block_id for record in spec_records)
    if expected == actual:
        return

    missing = sorted((expected - actual).elements())
    unexpected = sorted((actual - expected).elements())
    details: list[str] = []
    if missing:
        details.append(f"missing detail IDs: {', '.join(missing[:5])}")
    if unexpected:
        details.append(f"unexpected detail IDs: {', '.join(unexpected[:5])}")
    raise ValueError("M8 summary/detail inputs are inconsistent (" + "; ".join(details) + ")")


def _set_widths(ws, widths: Iterable[int]) -> None:
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(index)].width = width


def _style_title(ws, title: str, end_column: int) -> None:
    ws.cell(row=2, column=1, value=title)
    ws.cell(row=2, column=1).font = Font(name="Arial", bold=True, size=15, color=NAVY)
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=end_column)
    ws.row_dimensions[2].height = 24
    for cell in ws[3][:end_column]:
        cell.border = Border(bottom=Side(style="medium", color=BLUE))


def _style_header(ws, row: int, column_count: int) -> None:
    for column in range(1, column_count + 1):
        cell = ws.cell(row=row, column=column)
        cell.font = HEADER_FONT
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def _style_body(ws, start_row: int, end_row: int, column_count: int) -> None:
    for row in range(start_row, end_row + 1):
        for column in range(1, column_count + 1):
            cell = ws.cell(row=row, column=column)
            cell.font = BODY_FONT
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.border = BOTTOM_BORDER


def _write_document_summary(
    ws,
    manifest: list[dict],
    classifications: list[ClassificationResult],
) -> None:
    _style_title(ws, "문서별 전장 요구사항 요약", len(SUMMARY_HEADERS))
    for column, header in enumerate(SUMMARY_HEADERS, start=1):
        ws.cell(row=5, column=column, value=header)
    _style_header(ws, 5, len(SUMMARY_HEADERS))

    by_doc: dict[str, list[ClassificationResult]] = defaultdict(list)
    for item in classifications:
        by_doc[item.doc_id].append(item)

    first_data_row = 6
    for row, entry in enumerate(manifest, start=first_data_row):
        doc_id = str(entry["doc_id"])
        items = by_doc.get(doc_id, [])
        total_requirements = sum(item.is_requirement for item in items)
        ee_requirements = sum(item.is_requirement and item.is_ee_related for item in items)
        is_ee_document = any(item.is_ee_related for item in items)
        values = [
            doc_id,
            Path(entry["path"]).name,
            str(entry["doc_type"]).upper(),
            "관련" if is_ee_document else "비관련",
            total_requirements,
            ee_requirements,
            ee_requirements / total_requirements if total_requirements else 0,
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=value)

    last_data_row = first_data_row + len(manifest) - 1
    total_row = last_data_row + 1
    ws.cell(row=total_row, column=1, value="합계")
    ws.cell(row=total_row, column=5, value=f"=SUM(E{first_data_row}:E{last_data_row})")
    ws.cell(row=total_row, column=6, value=f"=SUM(F{first_data_row}:F{last_data_row})")
    ws.cell(row=total_row, column=7, value=f'=IF(E{total_row}=0,0,F{total_row}/E{total_row})')
    _style_body(ws, first_data_row, total_row, len(SUMMARY_HEADERS))
    for cell in ws[total_row][: len(SUMMARY_HEADERS)]:
        cell.font = Font(name="Arial", size=10, bold=True, color=NAVY)
        cell.fill = PatternFill("solid", fgColor=LIGHT_BLUE)
    for row in range(first_data_row, total_row + 1):
        ws.cell(row=row, column=4).alignment = Alignment(horizontal="center", vertical="center")
        for column in (5, 6, 7):
            ws.cell(row=row, column=column).alignment = Alignment(horizontal="right", vertical="center")
        ws.cell(row=row, column=7).number_format = "0.0%"

    ws.conditional_formatting.add(
        f"D{first_data_row}:D{last_data_row}",
        FormulaRule(formula=[f'D{first_data_row}="관련"'], fill=PatternFill("solid", fgColor=LIGHT_GREEN)),
    )
    ws.conditional_formatting.add(
        f"D{first_data_row}:D{last_data_row}",
        FormulaRule(formula=[f'D{first_data_row}="비관련"'], fill=PatternFill("solid", fgColor=LIGHT_RED)),
    )
    _set_widths(ws, [28, 38, 12, 18, 18, 24, 22])
    ws.freeze_panes = "A6"
    ws.auto_filter.ref = f"A5:G{last_data_row}"
    ws.sheet_view.showGridLines = False


def _source_note(record: SpecRecord) -> str:
    source = record.source_file
    if record.page_number is not None:
        source = f"{source}, p./slide {record.page_number}"
    return "\n".join(part for part in (record.notes.strip(), source.strip()) if part)


def _write_requirement_detail(ws, records: list[SpecRecord]) -> None:
    _style_title(ws, "전장품 관련 요구사항", len(DETAIL_HEADERS))
    for column, header in enumerate(DETAIL_HEADERS, start=1):
        ws.cell(row=5, column=column, value=header)
    _style_header(ws, 5, len(DETAIL_HEADERS))

    first_data_row = 6
    for row, record in enumerate(records, start=first_data_row):
        values = [
            record.doc_id,
            record.cur_id or record.record_id or record.block_id,
            record.english_text,
            record.korean_translation,
            _source_note(record),
            CATEGORY_KO.get(record.ee_category or "", record.ee_category or "기타"),
        ]
        for column, value in enumerate(values, start=1):
            ws.cell(row=row, column=column, value=value)
    last_data_row = max(5, first_data_row + len(records) - 1)
    if records:
        _style_body(ws, first_data_row, last_data_row, len(DETAIL_HEADERS))
        for row in range(first_data_row, last_data_row + 1):
            longest_text = max(
                len(str(ws.cell(row=row, column=3).value or "")),
                len(str(ws.cell(row=row, column=4).value or "")),
                len(str(ws.cell(row=row, column=5).value or "")),
            )
            estimated_lines = max(3, ceil(longest_text / 95))
            ws.row_dimensions[row].height = min(300, 12 + estimated_lines * 15)
            ws.cell(row=row, column=6).alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    _set_widths(ws, [28, 24, 76, 76, 48, 24])
    ws.freeze_panes = "C6"
    ws.auto_filter.ref = f"A5:F{last_data_row}"
    ws.sheet_view.showGridLines = False


def generate_report(
    config: PipelineConfig,
    manifest: list[dict],
    prescreen_results: list[FileScreenResult],
    classifications: list[ClassificationResult],
    spec_records: list[SpecRecord],
) -> Path:
    """Generate the requested document summary and E/E requirement detail sheets."""
    del prescreen_results  # Kept in the signature for pipeline compatibility.
    _validate_report_inputs(classifications, spec_records)
    config.prepare_run_dir()
    workbook = Workbook()
    summary_sheet = workbook.active
    summary_sheet.title = "문서 요약"
    _write_document_summary(summary_sheet, manifest, classifications)
    _write_requirement_detail(workbook.create_sheet("전장 요구사항"), spec_records)

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    path = config.run_dir / f"final_report_{timestamp}.xlsx"
    workbook.save(path)
    print(f"[M8] Report generated → {path}")
    return path


def run(
    config: PipelineConfig,
    manifest: list[dict],
    prescreen_results: list[FileScreenResult],
    classifications: list[ClassificationResult],
    spec_records: list[SpecRecord],
) -> Path:
    """Execute M8."""
    return generate_report(config, manifest, prescreen_results, classifications, spec_records)
